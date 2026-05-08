"""
Build Midterm Take-home workflow reference document (Chinese).
Topic: Multivariate analysis comparing L. carteri vs L. torrens biting flies.
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# ---- 全局样式 (中英文字体) ----
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

# 页边距
for section in doc.sections:
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

def set_cn_font(run, font='宋体', size=11, bold=False, color=None):
    run.font.name = 'Calibri'
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.rPr
    if rPr is None:
        rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), font)

def add_heading(text, level=1):
    sizes = {1: 18, 2: 14, 3: 12}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_cn_font(r, '黑体', sizes.get(level, 11), bold=True, color=RGBColor(0x1F, 0x3A, 0x5F))
    return p

def add_para(text, bold=False, indent=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.4
    if indent:
        p.paragraph_format.first_line_indent = Pt(22)
    r = p.add_run(text)
    set_cn_font(r, '宋体', 11, bold=bold)
    return p

def add_bullet(text, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.6 + level * 0.6)
    r = p.add_run(text)
    set_cn_font(r, '宋体', 11)
    return p

def add_code(code, label=None):
    if label:
        p = doc.add_paragraph()
        r = p.add_run(label)
        set_cn_font(r, '黑体', 11, bold=True, color=RGBColor(0x33, 0x33, 0x33))
    for line in code.splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(line if line else ' ')
        r.font.name = 'Consolas'
        r.font.size = Pt(9.5)
        rPr = r._element.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            rPr.append(rFonts)
        rFonts.set(qn('w:ascii'), 'Consolas')
        rFonts.set(qn('w:hAnsi'), 'Consolas')
        rFonts.set(qn('w:eastAsia'), 'Consolas')

def add_table(rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = 'Light Grid Accent 1'
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = t.rows[i].cells[j]
            c.text = ''
            p = c.paragraphs[0]
            r = p.add_run(str(cell))
            set_cn_font(r, '宋体', 10, bold=(i == 0))
    return t

# =========================
# 标题
# =========================
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run('Statistics 232C  Midterm Take-home  完整分析流程参考')
set_cn_font(r, '黑体', 20, bold=True, color=RGBColor(0x1F, 0x3A, 0x5F))

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run('两种 Leptoconops 咬人蝇 ( L. torrens vs L. carteri ) 的多元比较')
set_cn_font(r, '宋体', 12)

doc.add_paragraph()

# =========================
# 1. 任务背景 & 数据
# =========================
add_heading('一、任务背景与数据说明', 1)
add_para('两种咬人蝇 L. torrens 与 L. carteri 在外形上极为相似，长期被认为是同一物种；但生物学差异（如新出生蝇的性别比、咬人习性）后被发现。本次作业基于 7 个形态学测量变量，判断两种蝇是否在多元意义下存在差异。', indent=True)

add_heading('1.1 数据 (flies.xls)', 2)
add_table([
    ['列', '变量', '英文表头', '说明'],
    ['1', 'X1', 'W-len', 'Wing length 翅长'],
    ['2', 'X2', 'W-wid', 'Wing width 翅宽'],
    ['3', 'X3', '3rd P-len', '3rd palp length 第三触须节长'],
    ['4', 'X4', '3rd P-wid', '3rd palp width 第三触须节宽'],
    ['5', 'X5', '4th P-len', '4th palp length 第四触须节长'],
    ['6', 'X6', 'Antl seg12', '第 12 节触角长'],
    ['7', 'X7', 'Antl seg13', '第 13 节触角长'],
    ['8', 'Y',  'Species',   '0 = L. torrens, 1 = L. carteri'],
])
add_para('样本量：n = 70（每种 35 只），变量数 p = 7。两组样本量相同便于后续协方差讨论与变换。', indent=True)

add_heading('1.2 题目要求 (要点)', 2)
add_bullet('图形化检验各变量及联合分布的正态性')
add_bullet('如有需要，进行变换 (Box-Cox / log 等)')
add_bullet('检测多元离群点')
add_bullet('对均值差 μ₁ − μ₂ 给出合适的同时置信区间')
add_bullet('给出方法、结果摘要与最终结论；附 R / Python 代码')

# =========================
# 2. 总体流程
# =========================
add_heading('二、整体分析流程 (一图概览)', 1)
flow = (
    '【Step 0】 读入数据 → 拆成两组 X(1)、X(2)\n'
    '   ↓\n'
    '【Step 1】 描述性统计：均值向量、协方差矩阵、相关矩阵、箱线图\n'
    '   ↓\n'
    '【Step 2】 单变量正态性：直方图 + Q-Q 图 + Shapiro-Wilk\n'
    '   ↓\n'
    '【Step 3】 多元正态性：χ²-Q-Q 图 (Mahalanobis D²) + Mardia 检验\n'
    '   ↓\n'
    '【Step 4】 离群点：Mahalanobis 距离阈值 χ²₇,0.975\n'
    '   ↓\n'
    '【Step 5】 必要时变换 (log / Box-Cox) 重做 Step 2-4\n'
    '   ↓\n'
    '【Step 6】 协方差矩阵齐性：Box M 检验\n'
    '   ↓\n'
    '【Step 7】 两组均值检验：Hotelling 两样本 T²\n'
    '   ↓\n'
    '【Step 8】 同时置信区间 (T² 同时 + Bonferroni)\n'
    '   ↓\n'
    '【Step 9】 结果整合 → 报告 + 代码附录'
)
add_code(flow, label='整体路线图')

# =========================
# 3. 各步骤详细方法
# =========================
add_heading('三、各步骤详细方法', 1)

# Step 0
add_heading('Step 0  数据读取与初步探查', 2)
add_para('• 用 R 的 readxl::read_excel() 或 Python 的 pandas.read_excel() 读取 flies.xls；')
add_para('• 检查缺失值、数据类型；将 Species 转为因子；')
add_para('• 把数据切成两个矩阵 X1 (L. torrens, n₁=35)、X2 (L. carteri, n₂=35)。')

# Step 1
add_heading('Step 1  描述性统计与可视化', 2)
add_bullet('对每组分别计算样本均值 x̄_g 与样本协方差 S_g (g = 1,2)')
add_bullet('箱线图 (boxplot by species)：直观比较各变量在两组的位置 / 分散')
add_bullet('成对散点图 (pairs plot)：检查线性、相关、离群点')
add_bullet('相关矩阵热力图：判断变量间共线性，决定是否需要降维')

# Step 2
add_heading('Step 2  单变量正态性检验', 2)
add_para('对 7 个变量，分别在 L. torrens / L. carteri 子样本里：', indent=True)
add_bullet('画 直方图 + 拟合密度曲线')
add_bullet('画 Q-Q 图 (qqnorm + qqline)')
add_bullet('做 Shapiro-Wilk 检验 (n=35 时该检验适用)')
add_para('若多数变量 Q-Q 图近似直线 + Shapiro p>0.05，可认为单变量近似正态，否则记录偏度/重尾迹象。', indent=True)

# Step 3
add_heading('Step 3  多元正态性检验', 2)
add_para('单变量正态是必要不是充分条件，必须做联合正态检验：', indent=True)
add_bullet('对每组计算 Mahalanobis 距离平方 d²_i = (x_i − x̄)ᵀ S⁻¹ (x_i − x̄)')
add_bullet('画 χ²-Q-Q 图：横轴 χ²₇ 分位数，纵轴排序后的 d²_i；近似 45° 直线则支持多元正态')
add_bullet("Mardia's 偏度与峰度检验 (R 包 MVN::mvn) 给出 p-value")
add_bullet('Henze-Zirkler 检验作为补充 (同样在 MVN 包内)')

# Step 4
add_heading('Step 4  离群点检测', 2)
add_bullet('用 Step 3 的 d²_i，与 χ²_{7,0.975} ≈ 16.01 比较；超过阈值的样本作为候选离群点')
add_bullet('在成对散点图 / Q-Q 图上标记，看是否反复出现')
add_bullet('谨慎处理：若一只蝇在多个变量都极端，应怀疑录入错误或亚种；不要随意删除，需在报告中讨论')

# Step 5
add_heading('Step 5  变换 (Transformation)', 2)
add_para('当 Step 2-3 提示偏态或重尾时：', indent=True)
add_bullet('形态学测量常用 自然对数 log(X) 或 Box-Cox 变换 (R: MASS::boxcox 或 forecast::BoxCox)')
add_bullet('对每个变量分别选择最优 λ；建议两组合并选 λ 以保持变量含义一致')
add_bullet('若 λ 接近 0 → 用 log；接近 0.5 → 平方根；接近 1 → 不变')
add_bullet('变换后重做 Step 2-3，确认正态性改善；若仍不达标，可改用稳健方法 (e.g. 稳健 Hotelling T²)')

# Step 6
add_heading('Step 6  协方差矩阵齐性 (Box\'s M)', 2)
add_para('两样本 Hotelling T² 的标准形式假设 Σ₁ = Σ₂。检验：', indent=True)
add_bullet("H₀: Σ₁ = Σ₂   vs   H₁: Σ₁ ≠ Σ₂")
add_bullet('用 R: biotools::boxM() 或 heplots::boxM()')
add_bullet('若 p > 0.05 → 用合并协方差 S_p；若 p < 0.05 → 改用 James / Yao 修正的 T²')

# Step 7
add_heading('Step 7  两样本 Hotelling T² 检验', 2)
add_code(
    'T² = (n₁n₂)/(n₁+n₂) · (x̄₁ − x̄₂)ᵀ S_p⁻¹ (x̄₁ − x̄₂)\n'
    'F  = (n₁+n₂−p−1) / (p(n₁+n₂−2)) · T²   ~   F_{p, n₁+n₂−p−1}\n'
    '此处 n₁ = n₂ = 35, p = 7,  分子自由度 = 7, 分母自由度 = 62',
    label='统计量'
)
add_para('R 实现：', indent=True)
add_code(
    'library(ICSNP)\n'
    'HotellingsT2(X1, X2)\n'
    '# 或手写: 见附录 A',
)

# Step 8
add_heading('Step 8  均值差的同时置信区间', 2)
add_para('一旦 H₀ 被拒，需要找出究竟在哪些变量上不同。两个常用方案：', indent=True)
add_heading('8.1  T² 同时置信区间 (Roy-Bose / Scheffé 类)', 3)
add_code(
    '对任意线性组合 aᵀ(μ₁ − μ₂):\n'
    'aᵀ(x̄₁ − x̄₂) ± sqrt( c² · aᵀ S_p a · (1/n₁+1/n₂) )\n'
    '其中 c² = (n₁+n₂−2)p / (n₁+n₂−p−1) · F_{p, n₁+n₂−p−1, α}\n'
    '取 a = e_k (单位向量) 即得第 k 个分量的同时区间'
)
add_heading('8.2  Bonferroni 同时置信区间', 3)
add_code(
    '对每个 k = 1,…,p:\n'
    '(x̄₁_k − x̄₂_k) ± t_{n₁+n₂−2, α/(2p)} · sqrt( s_{p,kk} · (1/n₁+1/n₂) )\n'
    '当只关心 p 个分量比较且 p 不大时，Bonferroni 区间通常更短'
)
add_para('实务建议：两套都给，比较哪个更窄；最终结论与不为零的区间相对应即"显著不同的变量"。', indent=True)

# Step 9
add_heading('Step 9  结果整合与报告', 2)
add_bullet('把所有图、表与 p 值汇总成一张总表（变量 × 检验 × p 值 / CI）')
add_bullet('明确写出：是否需要变换；离群点处理；协方差是否齐；T² 是否拒绝 H₀；哪些变量主导差异')
add_bullet('代码作为 Appendix 附后')

# =========================
# 4. 推荐报告结构
# =========================
add_heading('四、推荐报告结构 (8–12 页)', 1)
add_table([
    ['章节', '建议内容', '建议页数'],
    ['封面', '题目 / 课程 / 两位作者姓名 / 提交日期', '1'],
    ['1. Introduction', '生物学背景；问题陈述；7 个变量含义', '0.5–1'],
    ['2. Data', '数据来源、样本量、初步描述统计、箱线图', '1'],
    ['3. Normality Check', 'Q-Q 图、Shapiro、χ²-Q-Q、Mardia 结果', '1.5'],
    ['4. Outliers & Transformation', 'Mahalanobis 离群点、log / Box-Cox 选择与对比图', '1'],
    ['5. Covariance Homogeneity', "Box's M 结果及对方法的影响", '0.5'],
    ['6. Hotelling T² Test', '检验值、p 值、结论', '0.5–1'],
    ['7. Simultaneous CIs', 'T² 与 Bonferroni 两套区间表 + 解释', '1.5'],
    ['8. Discussion & Conclusion', '哪些变量驱动差异；与文献对照；局限', '1'],
    ['Appendix', '完整 R / Python 代码 + sessionInfo()', '2–3'],
])

# =========================
# 5. 代码附录 (R)
# =========================
add_heading('五、附录 A：R 代码模板 (可直接跑)', 1)

code_R = r'''
# ============================================================
# Stat 232C Midterm — Multivariate analysis of Leptoconops flies
# ============================================================
library(readxl)        # 读 .xls
library(MVN)           # Mardia / Henze-Zirkler
library(biotools)      # boxM
library(ICSNP)         # HotellingsT2
library(MASS)          # boxcox
library(GGally)        # ggpairs
library(ggplot2)

# ---- 0. 读数据 ----
df <- read_excel("flies.xls")
names(df) <- c("W_len","W_wid","P3_len","P3_wid","P4_len",
               "Ant12","Ant13","Species")
df$Species <- factor(df$Species, levels = c(0,1),
                     labels = c("L.torrens","L.carteri"))
X  <- as.matrix(df[, 1:7])
g  <- df$Species
X1 <- X[g == "L.torrens", ]
X2 <- X[g == "L.carteri", ]

# ---- 1. 描述性 ----
summary(df)
aggregate(. ~ Species, data = df, mean)
ggpairs(df, columns = 1:7, aes(color = Species, alpha = 0.6))

# ---- 2. 单变量正态 ----
par(mfrow = c(2, 4))
for (k in 1:7) {
  qqnorm(X1[, k], main = paste(colnames(X)[k], "torrens"))
  qqline(X1[, k])
}
shapiro_p <- sapply(1:7, function(k) {
  c(torrens  = shapiro.test(X1[, k])$p.value,
    carteri  = shapiro.test(X2[, k])$p.value)
})
colnames(shapiro_p) <- colnames(X); shapiro_p

# ---- 3. 多元正态 ----
mvn(X1, mvnTest = "mardia")
mvn(X2, mvnTest = "mardia")
mvn(X1, mvnTest = "hz")        # Henze-Zirkler

# χ²-Q-Q 图
chiQQ <- function(M, ttl) {
  d2 <- mahalanobis(M, colMeans(M), cov(M))
  qchi <- qchisq(ppoints(nrow(M)), df = ncol(M))
  plot(qchi, sort(d2), main = ttl,
       xlab = expression(chi[7]^2~quantile),
       ylab = "ordered d^2"); abline(0, 1, col = "red")
}
par(mfrow = c(1, 2))
chiQQ(X1, "L. torrens"); chiQQ(X2, "L. carteri")

# ---- 4. 离群点 ----
d2_1 <- mahalanobis(X1, colMeans(X1), cov(X1))
d2_2 <- mahalanobis(X2, colMeans(X2), cov(X2))
cutoff <- qchisq(0.975, df = 7)        # ≈ 16.01
which(d2_1 > cutoff); which(d2_2 > cutoff)

# ---- 5. 变换 (示例: 对所有列尝试 Box-Cox) ----
lambda <- sapply(1:7, function(k) {
  bc <- boxcox(X[, k] ~ 1, plotit = FALSE)
  bc$x[which.max(bc$y)]
})
names(lambda) <- colnames(X); round(lambda, 2)
# 若 lambda 接近 0 → log；本题中常常 log 即可
Xlog <- log(X)

# ---- 6. 协方差齐性 ----
boxM(X, g)        # 用原始数据
boxM(Xlog, g)     # 用变换后数据

# ---- 7. Hotelling T² ----
ht <- HotellingsT2(X1, X2)        # 原始
ht_log <- HotellingsT2(log(X1), log(X2))
ht; ht_log

# ---- 8. 同时置信区间 ----
n1 <- nrow(X1); n2 <- nrow(X2); p <- 7
xbar1 <- colMeans(X1); xbar2 <- colMeans(X2)
Sp <- ((n1 - 1) * cov(X1) + (n2 - 1) * cov(X2)) / (n1 + n2 - 2)

alpha <- 0.05
c2 <- (n1 + n2 - 2) * p / (n1 + n2 - p - 1) *
      qf(1 - alpha, p, n1 + n2 - p - 1)

CI_T2 <- t(sapply(1:p, function(k) {
  diff <- xbar1[k] - xbar2[k]
  se   <- sqrt(c2 * Sp[k, k] * (1 / n1 + 1 / n2))
  c(diff - se, diff + se)
}))
rownames(CI_T2) <- colnames(X); CI_T2

t_bf <- qt(1 - alpha / (2 * p), n1 + n2 - 2)
CI_BF <- t(sapply(1:p, function(k) {
  diff <- xbar1[k] - xbar2[k]
  se   <- t_bf * sqrt(Sp[k, k] * (1 / n1 + 1 / n2))
  c(diff - se, diff + se)
}))
rownames(CI_BF) <- colnames(X); CI_BF

# ---- 9. 输出汇总 ----
result_tbl <- data.frame(
  Variable     = colnames(X),
  Mean_torrens = round(xbar1, 2),
  Mean_carteri = round(xbar2, 2),
  Diff         = round(xbar1 - xbar2, 2),
  T2_Lo        = round(CI_T2[, 1], 2),
  T2_Hi        = round(CI_T2[, 2], 2),
  BF_Lo        = round(CI_BF[, 1], 2),
  BF_Hi        = round(CI_BF[, 2], 2)
)
print(result_tbl)
'''
add_code(code_R.strip(), label='R 代码 (建议在 RStudio 中分块运行)')

# =========================
# 6. Python 替代实现
# =========================
add_heading('六、附录 B：Python 替代实现', 1)
code_py = r'''
import numpy as np, pandas as pd, scipy.stats as st
import matplotlib.pyplot as plt

df = pd.read_excel("flies.xls")
df.columns = ["W_len","W_wid","P3_len","P3_wid","P4_len",
              "Ant12","Ant13","Species"]
X1 = df[df.Species == 0].iloc[:, :7].values
X2 = df[df.Species == 1].iloc[:, :7].values
n1, n2 = len(X1), len(X2); p = 7

xbar1, xbar2 = X1.mean(0), X2.mean(0)
S1, S2 = np.cov(X1, rowvar=False), np.cov(X2, rowvar=False)
Sp = ((n1-1)*S1 + (n2-1)*S2) / (n1 + n2 - 2)

# Hotelling T^2
diff = xbar1 - xbar2
T2 = n1*n2/(n1+n2) * diff @ np.linalg.solve(Sp, diff)
F  = (n1+n2-p-1)/(p*(n1+n2-2)) * T2
pval = 1 - st.f.cdf(F, p, n1+n2-p-1)
print("T2 =", T2, " F =", F, " p =", pval)

# Mahalanobis 离群点
def mahal(M):
    mu, S = M.mean(0), np.cov(M, rowvar=False)
    inv = np.linalg.inv(S)
    return np.array([(x-mu) @ inv @ (x-mu) for x in M])
cut = st.chi2.ppf(0.975, p)
print("Outliers torrens:", np.where(mahal(X1) > cut)[0])
print("Outliers carteri:", np.where(mahal(X2) > cut)[0])

# 同时置信区间
c2 = (n1+n2-2)*p/(n1+n2-p-1) * st.f.ppf(0.95, p, n1+n2-p-1)
se_T2 = np.sqrt(c2 * np.diag(Sp) * (1/n1 + 1/n2))
se_BF = st.t.ppf(1 - 0.05/(2*p), n1+n2-2) * np.sqrt(np.diag(Sp)*(1/n1+1/n2))
out = pd.DataFrame({
    "Var":  df.columns[:7],
    "Diff": diff,
    "T2_Lo": diff - se_T2, "T2_Hi": diff + se_T2,
    "BF_Lo": diff - se_BF, "BF_Hi": diff + se_BF
}); print(out.round(2))
'''
add_code(code_py.strip(), label='Python (numpy + scipy)')

# =========================
# 7. 写报告时的 “避坑” 清单
# =========================
add_heading('七、写报告时的"避坑清单"', 1)
add_bullet('不要直接跳过正态性检验就上 Hotelling T²；老师会查 Q-Q 图和 Mardia')
add_bullet("Box's M 对正态偏离敏感；如果 M 检验显著但样本均衡 (n₁=n₂=35)，T² 仍较稳健，可注明并继续")
add_bullet('同时置信区间一定要写清楚整体置信水平 1−α，不要把每个变量的边际 CI 当成"同时" CI')
add_bullet('离群点处理要透明：保留 / 删除两版结果都跑一下做敏感性分析')
add_bullet('文末结论要用业务语言：到底两种蝇在哪些形态学指标上有差，差多少 (单位)')
add_bullet('代码 Appendix 附 sessionInfo() / pip freeze，方便复现')

doc.add_paragraph()
end = doc.add_paragraph()
end.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = end.add_run('— 完整流程文档 END —')
set_cn_font(r, '宋体', 10, color=RGBColor(0x88, 0x88, 0x88))

out_path = '/sessions/admiring-eager-cerf/mnt/outputs/Midterm_完整流程参考.docx'
doc.save(out_path)
print('Saved:', out_path)
